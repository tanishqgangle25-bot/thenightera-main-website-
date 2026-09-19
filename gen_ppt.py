import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# We want a dark theme, so we'll customize the slide backgrounds
def apply_dark_theme(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(29, 29, 31) # Dark grey/black matching thenightera vibe

def set_text_color(shape, r=255, g=255, b=255):
    if not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        for r_run in p.runs:
            r_run.font.color.rgb = RGBColor(r, g, b)

# Slide 1: Title
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "The Aroma White Audit"
title.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
title.text_frame.paragraphs[0].font.size = Pt(44)
title.text_frame.paragraphs[0].font.bold = True

subtitle.text = "Exposing the Brand Gap\n\nAudited by thenightera\n(Premium Brand & Web Architects)"
set_text_color(subtitle, 180, 180, 180)

# Slide 2: The Brutal Reality
slide_layout = prs.slide_layouts[1] # Title and Content
slide = prs.slides.add_slide(slide_layout)
apply_dark_theme(slide)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "The Brutal Reality (Current Stats)"
set_text_color(title, 255, 60, 60) # Red warning

tf = content.text_frame
tf.text = "Instagram Profile Data:"
p = tf.add_paragraph()
p.text = "• Followers: 158"
p.level = 1
p = tf.add_paragraph()
p.text = "• Following: 78"
p.level = 1
p = tf.add_paragraph()
p.text = "• Total Posts: 40"
p.level = 1

p = tf.add_paragraph()
p.text = "\nTHE 158 FOLLOWER RED FLAG:"
p.font.bold = True
p.level = 0
p = tf.add_paragraph()
p.text = "40 posts in and only 158 followers is a critical failure in content distribution. It means the content has ZERO organic reach. In a competitive market like Mahalaxmi Nagar, having under 200 followers makes a cafe look unappealing to new visitors."
p.level = 1
set_text_color(content, 230, 230, 230)


# Slide 3: The Core Mistakes
slide = prs.slides.add_slide(prs.slide_layouts[1])
apply_dark_theme(slide)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "The Core Mistakes (Why It's Failing)"
set_text_color(title, 255, 255, 255)

tf = content.text_frame
tf.text = "1. Treating Instagram Like a Menu Board"
p = tf.add_paragraph()
p.text = "Posting static photos of coffee cups with zero storytelling. People follow for the vibe and FOMO, not a static sandwich picture."
p.level = 1

p = tf.add_paragraph()
p.text = "2. Zero Viral Hook Strategy (The Video Deficit)"
p.level = 0
p = tf.add_paragraph()
p.text = "Generic Canva templates and basic phone videos do not trigger the Instagram algorithm. Without high-retention cinematic Reels, growth is impossible."
p.level = 1

p = tf.add_paragraph()
p.text = "3. No Premium Brand Identity"
p.level = 0
p = tf.add_paragraph()
p.text = "The grid lacks a cohesive color grading strategy, highlights are messy, and the bio doesn't convert profile visits into actual footfall."
p.level = 1
set_text_color(content, 230, 230, 230)

# Slide 4: Why they need thenightera
slide = prs.slides.add_slide(prs.slide_layouts[1])
apply_dark_theme(slide)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "The Solution: Why You Need thenightera"
set_text_color(title, 236, 230, 216) # Goldish accent

tf = content.text_frame
tf.text = "To beat cafes with massive budgets, you need a Visual Monopoly:"
p = tf.add_paragraph()
p.text = "1. Heavy VFX & Cinematic Reels"
p.font.bold = True
p.level = 1
p = tf.add_paragraph()
p.text = "Replacing boring statics with high-end, VFX-integrated video content that forces users to stop scrolling."
p.level = 2

p = tf.add_paragraph()
p.text = "2. The Premium Web Ecosystem"
p.font.bold = True
p.level = 1
p = tf.add_paragraph()
p.text = "Building a liquid-smooth, 3D interactive website. When a customer clicks the bio link, they enter a premium digital vault to book tables."
p.level = 2

p = tf.add_paragraph()
p.text = "3. Elite Aesthetic Overhaul"
p.font.bold = True
p.level = 1
p = tf.add_paragraph()
p.text = "Complete redesign of the grid with dark-mode aesthetics and strict color palettes to look like an international franchise."
p.level = 2
set_text_color(content, 230, 230, 230)

# Slide 5: ROI
slide = prs.slides.add_slide(prs.slide_layouts[1])
apply_dark_theme(slide)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "The Expected ROI"
set_text_color(title, 150, 255, 150) # Light green

tf = content.text_frame
tf.text = "Trajectory after hiring thenightera:"
p = tf.add_paragraph()
p.text = "Month 1: The Rebrand"
p.font.bold = True
p.level = 1
p = tf.add_paragraph()
p.text = "Complete grid cleanup. Followers shift from a dead 158 to a highly engaged local community via targeted Reel distribution."
p.level = 2

p = tf.add_paragraph()
p.text = "Month 2: The Hype"
p.font.bold = True
p.level = 1
p = tf.add_paragraph()
p.text = "Viral local reach. People in Indore will start sharing the cinematic reels in DMs."
p.level = 2

p = tf.add_paragraph()
p.text = "Month 3: The Footfall Shift"
p.font.bold = True
p.level = 1
p = tf.add_paragraph()
p.text = "Digital traffic converts to physical footfall. Higher bookings, higher ticket sizes (premium brand = premium prices), and a cult-like following."
p.level = 2
set_text_color(content, 230, 230, 230)

# Save presentation
prs.save('Cafe_Aroma_White_Pitch.pptx')
